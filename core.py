"""Independent, evidence-first procurement workspace. No upstream code reused."""
from __future__ import annotations
import hashlib
import json
import os
import re
import tempfile
import unicodedata
from datetime import datetime, timezone, timedelta
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

CHILE = ZoneInfo('America/Santiago')
GROUPS = {
    'bienestar y autocuidado': ['bienestar', 'autocuidado', 'mindfulness', 'yoga', 'estres', 'regulacion emocional'],
    'equipos y organizaciones': ['liderazgo', 'comunicacion', 'trabajo en equipo', 'clima laboral', 'clima organizacional', 'conflictos', 'coaching'],
    'teatro y expresión': ['teatro', 'narracion oral', 'expresion corporal', 'memoria', 'escritura terapeutica'],
    'diagnóstico e intervención': ['diagnostico organizacional', 'encuesta', 'intervencion grupal'],
}

def now():
    return datetime.now(timezone.utc).isoformat()

def norm(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value or '').lower()) if not unicodedata.combining(c))

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def safe_id(value):
    if not re.fullmatch(r'[A-Za-z0-9_-]{1,100}', value):
        raise ValueError('Identificador inválido')
    return value

def private_json(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    fd, temp = tempfile.mkstemp(dir=path.parent)
    try:
        with os.fdopen(fd, 'w') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        os.replace(temp, path)
        path.chmod(0o600)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)

def read_json(path, default=None):
    return json.loads(Path(path).read_text()) if Path(path).exists() else default

def date_info(raw, verified=False, checked_at=None):
    result = {'original': str(raw or ''), 'zona': 'America/Santiago', 'verificada': verified, 'verificada_en': checked_at}
    try:
        dt = datetime.fromisoformat(str(raw).replace('Z', '+00:00'))
        if dt.tzinfo is None:
            result['advertencia'] = 'Fecha sin zona horaria: confirmar en portal'
        else:
            result['conversion_chile'] = dt.astimezone(CHILE).isoformat()
    except (ValueError, TypeError):
        result['advertencia'] = 'Fecha ausente o no interpretable'
    if not verified:
        result['advertencia'] = 'Conversión del dato fuente; cierre pendiente de contraste con portal'
    return result

def affinity(record):
    text = norm(record.get('nombre', '') + ' ' + record.get('descripcion', ''))
    hits = {group: [word for word in words if re.search(r'(?<!\w)' + re.escape(word) + r'(?!\w)', text)] for group, words in GROUPS.items()}
    hits = {k: v for k, v in hits.items() if v}
    return {'nivel': 'ALTA' if len(hits) >= 2 else 'MEDIA' if hits else 'SIN COINCIDENCIAS',
            'motivos': hits, 'criterio': 'Coincidencias temáticas, no probabilidad de adjudicación', 'orden': len(hits)}

def category_filter(records, codes):
    """Apply requested UNSPSC prefixes to the entire input, never a software subset."""
    prefixes = [str(c) for c in codes]
    if not prefixes or any(not p.isdigit() for p in prefixes):
        raise ValueError('Indicar códigos UNSPSC numéricos')
    return [r for r in records if any(str(code).startswith(tuple(prefixes)) for code in r.get('codigos_unspsc', []))]

def import_excel(path):
    from openpyxl import load_workbook
    path = Path(path).expanduser().resolve()
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = iter(wb.active.values)
    headers = list(next(rows))
    required = ['ID', 'Nombre', 'Fecha de cierre', 'Monto Disponible', 'Organismo']
    if any(k not in headers for k in required):
        raise ValueError('Excel no reconocido: faltan columnas del portal')
    sha, imported, result, seen = digest(path), now(), [], set()
    for row_number, values in enumerate(rows, 2):
        row = dict(zip(headers, values))
        if not row.get('ID'):
            continue
        code = safe_id(str(row['ID']))
        if code in seen:
            raise ValueError('Código repetido en Excel: ' + code)
        seen.add(code)
        raw_amount = row['Monto Disponible']
        amount = None
        if isinstance(raw_amount, (float, int)):
            amount = str(raw_amount)
        elif re.fullmatch(r'\d{1,3}(\.\d{3})*|\d+', str(raw_amount or '')):
            amount = str(raw_amount).replace('.', '')
        record = {'codigo': code, 'tipo': 'compra_agil', 'nombre': str(row['Nombre'] or ''),
                  'organismo': row['Organismo'], 'cierre': date_info(row['Fecha de cierre']),
                  'presupuesto': {'original': raw_amount, 'monto': amount, 'moneda': 'NO CONFIRMADA', 'impuestos': 'DESCONOCIDOS'},
                  'estado_convocatoria': row.get('Estado Convocatoria'),
                  'url_portal': f'https://compra-agil.mercadopublico.cl/resumen-cotizacion/{code}',
                  'fuente': {'archivo': str(path), 'sha256': sha, 'hoja': wb.active.title, 'fila': row_number, 'importado_en': imported}}
        record['afinidad'] = affinity(record)
        result.append(record)
    wb.close()
    return result

def extract_document(path):
    path = Path(path).expanduser().resolve()
    suffix = path.suffix.lower()
    parts = []
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        for n, page in enumerate(reader.pages, 1):
            parts.append({'ubicacion': f'página {n}', 'texto': page.extract_text() or ''})
    elif suffix == '.docx':
        from docx import Document
        doc = Document(path)
        for n, p in enumerate(doc.paragraphs, 1):
            parts.append({'ubicacion': f'párrafo {n}', 'texto': p.text})
        for n, table in enumerate(doc.tables, 1):
            for row_n, row in enumerate(table.rows, 1):
                parts.append({'ubicacion': f'tabla {n}, fila {row_n}', 'texto': ' | '.join(c.text for c in row.cells)})
        for n, section in enumerate(doc.sections, 1):
            for label, area in [('encabezado', section.header), ('pie', section.footer)]:
                parts.append({'ubicacion': f'{label} sección {n}', 'texto': '\n'.join(p.text for p in area.paragraphs)})
    elif suffix in ('.txt', '.md', '.csv'):
        parts = [{'ubicacion': f'línea {n}', 'texto': line} for n, line in enumerate(path.read_text().splitlines(), 1)]
    else:
        raise ValueError('Formato no soportado: convertir DOC/imagen a PDF con texto o DOCX; no marcar leído')
    return {'archivo': str(path), 'sha256': digest(path), 'extraido_en': now(), 'partes': parts,
            'requiere_revision_visual': suffix in ('.pdf', '.docx'),
            'paginas_sin_texto': [p['ubicacion'] for p in parts if suffix == '.pdf' and not p['texto'].strip()],
            'lectura_confirmada': False}

def evidence_valid(evidence, documents):
    if not isinstance(evidence, dict):
        return False
    doc = documents.get(evidence.get('documento'))
    quote = evidence.get('cita', '').strip()
    if not doc or not quote or evidence.get('sha256') != doc['sha256']:
        return False
    if not Path(doc['archivo']).is_file() or digest(doc['archivo']) != doc['sha256']:
        return False
    return any(p['ubicacion'] == evidence.get('ubicacion') and quote in p['texto'] for p in doc['partes'])

def review(record, dossier, at=None):
    at = at or datetime.now(timezone.utc)
    docs = dossier.get('documentos', {})
    pending, excludes = [], []
    inventory = dossier.get('inventario', {})
    if not inventory.get('confirmado') or not inventory.get('fuente') or not inventory.get('revisado_en'):
        pending.append('Confirmar inventario completo de bases, anexos, respuestas y modificaciones en portal')
    else:
        try:
            age = at - datetime.fromisoformat(inventory['revisado_en'])
            if age < timedelta(0) or age > timedelta(hours=24):
                pending.append('Actualizar inventario del portal (más de 24 horas o fecha inválida)')
        except (ValueError, TypeError):
            pending.append('Fecha de revisión del inventario inválida')
    for name in inventory.get('esperados', []):
        if name not in docs:
            pending.append('Falta documento: ' + name)
    if not docs:
        pending.append('Faltan bases/documentos completos')
    for name, doc in docs.items():
        if not doc.get('lectura_confirmada') or doc.get('paginas_sin_texto'):
            pending.append('Revisar documento completo/visual: ' + name)
        if not Path(doc['archivo']).exists() or digest(doc['archivo']) != doc['sha256']:
            pending.append('Documento cambiado o ausente: ' + name)
    otec = dossier.get('otec', {})
    status = otec.get('estado', 'DESCONOCIDO')
    if status not in ('OBLIGATORIO', 'PUNTUABLE', 'NO_REQUERIDO') or not evidence_valid(otec.get('evidencia'), docs):
        status = 'DESCONOCIDO'
        pending.append('Confirmar exigencia OTEC desde las bases; ausencia de palabra no prueba exención')
    if status == 'OBLIGATORIO':
        excludes.append('Las bases exigen OTEC y Camino Blanco no es OTEC')
    requirements = dossier.get('requisitos', [])
    if not dossier.get('requisitos_revisados') or not requirements:
        pending.append('Revisar todos los requisitos de admisibilidad y ejecución')
    for item in requirements:
        if not evidence_valid(item.get('evidencia'), docs):
            pending.append('Falta respaldo en bases: ' + item.get('nombre', 'requisito'))
        elif item.get('excluyente') and item.get('cumple') is False:
            excludes.append(item['nombre'])
        elif item.get('cumple') is not True:
            pending.append('Verificar: ' + item.get('nombre', 'requisito'))
        elif not item.get('respaldo_proveedor'):
            pending.append('Falta antecedente del proveedor: ' + item.get('nombre', 'requisito'))
    close = dossier.get('cierre', {})
    try:
        dt = datetime.fromisoformat(close.get('valor', '').replace('Z', '+00:00'))
        checked = datetime.fromisoformat(close.get('revisado_en', ''))
        if not dt.tzinfo or not checked.tzinfo or not close.get('fuente') or not timedelta(0) <= at - checked <= timedelta(hours=24):
            raise ValueError()
        if at >= dt:
            excludes.append('Plazo de postulación cerrado según fuente verificada')
    except (ValueError, TypeError):
        pending.append('Confirmar cierre vigente y hora de Chile')
    budget = dossier.get('presupuesto', {})
    valid_budget = False
    try:
        amount = Decimal(str(budget.get('monto')))
        valid_budget = amount.is_finite() and amount >= 0 and bool(budget.get('moneda'))
    except Exception:
        pass
    if not valid_budget or budget.get('impuestos') not in ('INCLUIDOS', 'EXCLUIDOS', 'EXENTO') or not budget.get('fuente'):
        pending.append('Confirmar presupuesto, moneda e impuestos')
    return {'codigo': record['codigo'], 'afinidad': affinity(record), 'otec': status,
            'decision': 'DESCARTAR' if excludes else 'INVESTIGAR' if pending else 'POSTULAR',
            'exclusiones': excludes, 'pendientes': pending,
            'evaluado_en': at.isoformat(), 'alcance': 'POSTULAR significa preparar para revisión; nunca envío ni firma automáticos'}

class Workspace:
    def __init__(self, root=None):
        self.root = Path(root or os.getenv('CAMINO_MP_DATA', Path.home() / '.local/share/camino-blanco-mercado-publico/data')).expanduser()
        self.root.mkdir(parents=True, exist_ok=True, mode=0o700)
        self.root.chmod(0o700)

    def path(self, code):
        return self.root / 'expedientes' / (safe_id(code) + '.json')

    def dossier(self, code):
        return read_json(self.path(code), {'codigo': code, 'documentos': {}})

    def records(self):
        return read_json(self.root / 'oportunidades.json', [])

    def merge(self, records):
        current = {r['codigo']: r for r in self.records()}
        for r in records:
            old = current.get(r['codigo'])
            if old and old != r:
                private_json(self.root / 'historial' / (safe_id(r['codigo']) + '-' + hashlib.sha256(json.dumps(old, sort_keys=True).encode()).hexdigest() + '.json'), old)
                dossier = self.dossier(r['codigo'])
                dossier['inventario'] = {}
                dossier['cierre'] = {}
                private_json(self.path(r['codigo']), dossier)
            current[r['codigo']] = r
        private_json(self.root / 'oportunidades.json', list(current.values()))
        return len(records)

    def list(self, query=''):
        rows = [r for r in self.records() if norm(query) in norm(json.dumps(r, ensure_ascii=False))]
        return sorted([dict(r, evaluacion=review(r, self.dossier(r['codigo'])), observaciones_portal=self.dossier(r['codigo']).get('observaciones_portal')) for r in rows], key=lambda r: (-r['evaluacion']['afinidad']['orden'], r['codigo']))
